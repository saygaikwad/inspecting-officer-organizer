import streamlit as st
import pandas as pd
from io import BytesIO
from openpyxl.styles import Border, Side, Font, Alignment
from openpyxl.utils import get_column_letter


# =========================================================
# PAGE CONFIGURATION
# =========================================================

st.set_page_config(
    page_title="Inspecting Officer Organizer | SSC",
    layout="wide"
)


# =========================================================
# CUSTOM CSS
# =========================================================

st.markdown("""
<style>

.nic-header {
    background: #7b1e1e;
    color: white;
    padding: 16px;
    font-size: 26px;
    font-weight: 700;
}

.nic-sub {
    background: white;
    padding: 10px;
    border-bottom: 3px solid #7b1e1e;
}

.nic-card {
    background: white;
    padding: 18px;
    border: 1px solid #ccc;
    border-radius: 8px;
    margin-bottom: 20px;
}

.nic-title {
    font-size: 18px;
    font-weight: bold;
    color: #7b1e1e;
    border-bottom: 2px solid #7b1e1e;
    padding-bottom: 6px;
    margin-bottom: 12px;
}

</style>
""", unsafe_allow_html=True)


# =========================================================
# HEADER
# =========================================================

st.markdown(
    '<div class="nic-header">Inspecting Officer Organizer</div>',
    unsafe_allow_html=True
)

st.markdown(
    '<div class="nic-sub">Staff Selection Commission | Examination Duty Management System</div>',
    unsafe_allow_html=True
)


# =========================================================
# EXAMINATION OPTIONS
# =========================================================

EXAMS = [
    "Combined Graduate Level Examination",
    "Combined Higher Secondary Examination",
    "Junior Engineer Examination",
    "Selection Post Examination"
]


# =========================================================
# SIDEBAR
# =========================================================

with st.sidebar:

    st.markdown("### 🛠 Control Panel")

    selected_exam = st.selectbox(
        "Examination",
        EXAMS
    )

    exam_year = st.selectbox(
        "Year",
        range(2025, 2037)
    )

    custom_exam_name = st.text_input(
        "Custom Examination Name"
    )

    st.markdown("---")

    uploaded_file = st.file_uploader(
        "Upload Inspecting Officers Master Excel",
        type=["xlsx"]
    )


# =========================================================
# CHECK FILE UPLOAD
# =========================================================

if uploaded_file is None:

    st.info("Please upload Excel file.")

    st.stop()


# =========================================================
# EXAMINATION NAME
# =========================================================

final_exam_name = (
    custom_exam_name.strip()
    if custom_exam_name.strip()
    else selected_exam
)


# =========================================================
# READ EXCEL
# =========================================================

try:

    df = pd.read_excel(
        uploaded_file,
        dtype=str
    )

except Exception as e:

    st.error(f"Error reading Excel: {e}")

    st.stop()


# =========================================================
# CLEAN COLUMN NAMES
# =========================================================

df.columns = (
    df.columns
    .astype(str)
    .str.strip()
)


# =========================================================
# COLUMN NAMES
# =========================================================

NAME_COL = "Name of Inspecting Officer"
GROUP_COL = "GROUP 'A' & 'B'"
STATUS_COL = "status"
BANK_COL = "Name of Bank & Branch"
ACC_COL = "Account No"
IFSC_COL = "IFSC CODE"
CITY_COL = "CITY"


# =========================================================
# REQUIRED COLUMNS
# =========================================================

required_cols = [
    NAME_COL,
    GROUP_COL,
    STATUS_COL,
    BANK_COL,
    ACC_COL,
    IFSC_COL,
    CITY_COL
]

missing = [
    column
    for column in required_cols
    if column not in df.columns
]

if missing:

    st.error(
        f"Missing Columns: {missing}"
    )

    st.stop()


# =========================================================
# CLEAN CITY COLUMN
# =========================================================

df[CITY_COL] = (
    df[CITY_COL]
    .fillna("")
    .astype(str)
    .str.strip()
)


# =========================================================
# CLEAN OFFICER NAME COLUMN
# =========================================================

df[NAME_COL] = (
    df[NAME_COL]
    .fillna("")
    .astype(str)
    .str.strip()
)


# =========================================================
# CLEAN STATUS COLUMN
# =========================================================

df[STATUS_COL] = (
    df[STATUS_COL]
    .fillna("")
    .astype(str)
    .str.strip()
    .str.lower()
)


# =========================================================
# CLEAN GROUP COLUMN
# =========================================================

df[GROUP_COL] = (
    df[GROUP_COL]
    .fillna("")
    .astype(str)
    .str.strip()
    .str.upper()
)


# =========================================================
# REMOVE COMPLETELY EMPTY ROWS
# =========================================================

df = df.dropna(how="all")


# =========================================================
# REMOVE ROWS WITHOUT OFFICER NAME
# =========================================================

df = df[
    df[NAME_COL] != ""
]


# =========================================================
# RESET INDEX
# =========================================================

df = df.reset_index(drop=True)


# =========================================================
# REMUNERATION FUNCTION
# =========================================================

def remuneration(
    group,
    full_days,
    single_days,
    status
):

    full_rate = 0
    single_rate = 0

    if status == "retired":

        if group == "A":

            full_rate = 2000
            single_rate = 1200

        elif group == "B":

            full_rate = 1500
            single_rate = 900

    elif status == "working":

        if group == "A":

            full_rate = 1200
            single_rate = 1000

        elif group == "B":

            full_rate = 900
            single_rate = 750

    return (
        full_days * full_rate
    ) + (
        single_days * single_rate
    )


# =========================================================
# DISTANCE ALLOWANCE FUNCTION
# =========================================================

def distance_allowance(km):

    if km <= 20:

        return 300

    elif km <= 50:

        return 500

    return 750


# =========================================================
# CREATE FORMATTED EXCEL FILE
# =========================================================

def create_bordered_excel(
    df,
    exam_name,
    exam_year
):

    output = BytesIO()

    with pd.ExcelWriter(
        output,
        engine="openpyxl"
    ) as writer:

        start_row = 3

        df.to_excel(
            writer,
            index=False,
            sheet_name="Final Summary",
            startrow=start_row
        )

        ws = writer.sheets["Final Summary"]

        # -------------------------------------------------
        # TITLE
        # -------------------------------------------------

        ws.merge_cells(
            start_row=1,
            start_column=1,
            end_row=1,
            end_column=df.shape[1]
        )

        title = ws.cell(1, 1)

        title.value = (
            f"{exam_name} - {exam_year}"
        )

        title.font = Font(
            size=14,
            bold=True
        )

        title.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

        ws.row_dimensions[1].height = 25

        # -------------------------------------------------
        # BORDER
        # -------------------------------------------------

        thin = Side(
            style="thin"
        )

        border = Border(
            left=thin,
            right=thin,
            top=thin,
            bottom=thin
        )

        # -------------------------------------------------
        # FORMAT CELLS
        # -------------------------------------------------

        for row in ws.iter_rows():

            for cell in row:

                cell.border = border

                cell.alignment = Alignment(
                    horizontal="center",
                    vertical="center"
                )

                if cell.row == start_row + 1:

                    cell.font = Font(
                        bold=True
                    )

        # -------------------------------------------------
        # COLUMN WIDTH
        # -------------------------------------------------

        for col in range(
            1,
            ws.max_column + 1
        ):

            column_letter = get_column_letter(col)

            max_length = 0

            for cell in ws[column_letter]:

                if cell.value is not None:

                    max_length = max(
                        max_length,
                        len(str(cell.value))
                    )

            ws.column_dimensions[
                column_letter
            ].width = min(
                max(max_length + 3, 15),
                35
            )

        # -------------------------------------------------
        # HEADER FORMATTING
        # -------------------------------------------------

        header_row = start_row + 1

        for cell in ws[header_row]:

            cell.font = Font(
                bold=True,
                color="FFFFFF"
            )

            cell.fill = __import__(
                "openpyxl"
            ).styles.PatternFill(
                fill_type="solid",
                fgColor="7B1E1E"
            )

            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True
            )

        ws.freeze_panes = "A5"

    output.seek(0)

    return output


# =========================================================
# SESSION STATE
# =========================================================

if "assignments" not in st.session_state:

    st.session_state.assignments = []


# =========================================================
# CITY & OFFICER SELECTION
# =========================================================

st.markdown(
    '<div class="nic-card">',
    unsafe_allow_html=True
)

st.markdown(
    '<div class="nic-title">Select City & Inspecting Officers</div>',
    unsafe_allow_html=True
)


# =========================================================
# CITY LIST
# =========================================================

cities = sorted(
    df[CITY_COL]
    .loc[lambda x: x != ""]
    .unique()
)


# =========================================================
# NO CITY FOUND
# =========================================================

if not cities:

    st.error(
        "No valid cities were found in the uploaded Excel file."
    )

    st.stop()


# =========================================================
# CITY SELECTBOX
# =========================================================

selected_city = st.selectbox(
    "City",
    cities
)


# =========================================================
# FILTER CITY DATA
# =========================================================

city_df = df[
    df[CITY_COL] == selected_city
].copy()


# =========================================================
# OFFICER LIST
# =========================================================

officers = sorted(
    city_df[NAME_COL]
    .loc[lambda x: x != ""]
    .unique()
)


# =========================================================
# OFFICER MULTISELECT
# =========================================================

selected_officers = st.multiselect(
    "Inspecting Officers",
    officers
)


st.markdown(
    "</div>",
    unsafe_allow_html=True
)


# =========================================================
# OFFICER-WISE ASSIGNMENT
# =========================================================

if selected_officers:

    st.markdown(
        '<div class="nic-card">',
        unsafe_allow_html=True
    )

    st.markdown(
        '<div class="nic-title">Officer-wise Assignment</div>',
        unsafe_allow_html=True
    )

    # -----------------------------------------------------
    # LOOP THROUGH SELECTED OFFICERS
    # -----------------------------------------------------

    for officer in selected_officers:

        officer_rows = city_df[
            city_df[NAME_COL] == officer
        ]

        if officer_rows.empty:

            st.warning(
                f"No data found for {officer}."
            )

            continue

        row = officer_rows.iloc[0]

        # -------------------------------------------------
        # GROUP
        # -------------------------------------------------

        group = str(
            row[GROUP_COL]
        ).strip().upper()

        # -------------------------------------------------
        # STATUS
        # -------------------------------------------------

        status = str(
            row[STATUS_COL]
        ).strip().lower()

        display_status = (
            status.capitalize()
            if status
            else "Unknown"
        )

        # -------------------------------------------------
        # EXPANDER
        # -------------------------------------------------

        with st.expander(
            f"👤 {officer} | Group {group} | {display_status}"
        ):

            # ---------------------------------------------
            # DISTANCE
            # ---------------------------------------------

            distance = st.number_input(
                "Distance from HQ (KM)",
                min_value=0,
                value=0,
                step=1,
                key=f"dist_{selected_city}_{officer}"
            )

            # ---------------------------------------------
            # FULL SHIFT DAYS
            # ---------------------------------------------

            full_days = st.number_input(
                "Full Shift Days",
                min_value=0,
                value=0,
                step=1,
                key=f"full_{selected_city}_{officer}"
            )

            # ---------------------------------------------
            # SINGLE SHIFT DAYS
            # ---------------------------------------------

            single_days = st.number_input(
                "Single Shift Days",
                min_value=0,
                value=0,
                step=1,
                key=f"single_{selected_city}_{officer}"
            )

            # ---------------------------------------------
            # ADD / UPDATE
            # ---------------------------------------------

            if st.button(
                "Add / Update",
                key=f"add_{selected_city}_{officer}"
            ):

                # -----------------------------------------
                # REMOVE OLD ENTRY FOR SAME OFFICER
                # -----------------------------------------

                st.session_state.assignments = [
                    a
                    for a in st.session_state.assignments
                    if not (
                        a["Name of Inspecting Officer"] == officer
                        and a["City"] == selected_city
                    )
                ]

                # -----------------------------------------
                # DISTANCE ALLOWANCE
                # -----------------------------------------

                da = distance_allowance(
                    distance
                )

                # -----------------------------------------
                # TOTAL DAYS
                # -----------------------------------------

                total_days = (
                    full_days + single_days
                )

                # -----------------------------------------
                # REMUNERATION
                # -----------------------------------------

                remuneration_amount = remuneration(
                    group,
                    full_days,
                    single_days,
                    status
                )

                # -----------------------------------------
                # ADD ASSIGNMENT
                # -----------------------------------------

                st.session_state.assignments.append({

                    "Name of Inspecting Officer": officer,

                    "City": row[CITY_COL],

                    "Group": group,

                    "Status": display_status,

                    "Bank Name & Branch": row[BANK_COL],

                    "Account No": row[ACC_COL],

                    "IFSC Code": row[IFSC_COL],

                    "Distance (KM)": distance,

                    "Full Shift Days": full_days,

                    "Single Shift Days": single_days,

                    "Total Days": total_days,

                    "Remuneration": remuneration_amount,

                    "Distance Allowance": (
                        da * total_days
                    )

                })

                st.success(
                    f"{officer} updated successfully."
                )

    st.markdown(
        "</div>",
        unsafe_allow_html=True
    )


# =========================================================
# FINAL SUMMARY
# =========================================================

if st.session_state.assignments:

    summary = pd.DataFrame(
        st.session_state.assignments
    )

    # -----------------------------------------------------
    # GRAND TOTAL
    # -----------------------------------------------------

    summary["Grand Total"] = (
        summary["Remuneration"]
        + summary["Distance Allowance"]
    )

    # -----------------------------------------------------
    # SERIAL NUMBER
    # -----------------------------------------------------

    summary.insert(
        0,
        "Sl. No.",
        range(
            1,
            len(summary) + 1
        )
    )

    # -----------------------------------------------------
    # COLUMN ORDER
    # -----------------------------------------------------

    summary = summary[
        [
            "Sl. No.",
            "Name of Inspecting Officer",
            "City",
            "Group",
            "Status",
            "Bank Name & Branch",
            "Account No",
            "IFSC Code",
            "Distance (KM)",
            "Full Shift Days",
            "Single Shift Days",
            "Total Days",
            "Remuneration",
            "Distance Allowance",
            "Grand Total"
        ]
    ]

    # -----------------------------------------------------
    # DISPLAY SUMMARY
    # -----------------------------------------------------

    st.markdown(
        "### 📊 Final Summary"
    )

    st.dataframe(
        summary,
        use_container_width=True,
        hide_index=True
    )

    # -----------------------------------------------------
    # CREATE EXCEL
    # -----------------------------------------------------

    excel_file = create_bordered_excel(
        summary,
        final_exam_name,
        exam_year
    )

    # -----------------------------------------------------
    # DOWNLOAD BUTTON
    # -----------------------------------------------------

    safe_exam_name = "".join(
        c if c.isalnum() or c in (" ", "-", "_") else "_"
        for c in final_exam_name
    ).strip()

    st.download_button(
        "⬇ Download Final Summary (Excel)",
        data=excel_file,
        file_name=(
            f"{safe_exam_name}_"
            f"{exam_year}_"
            f"Final_Summary.xlsx"
        ),
        mime=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )
    )

else:

    st.info(
        "Select an inspecting officer, enter duty details, "
        "and click Add / Update to generate the final summary."
    )
